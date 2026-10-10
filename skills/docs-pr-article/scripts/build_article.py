#!/usr/bin/env python3
"""Build a Chinese PR article as .docx and .pdf from one plain-text source.

The deliverables are the .docx (pasted into the 公众号 editor) and the .pdf
(previewed remotely). The source is a build input, never a deliverable, so
it lives in a gitignored build directory beside the crops it references.

Source format, one directive or paragraph per line, blank lines ignored:

  @title     主标题
  @subtitle  副标题                     rendered as a "——副标题" line
  @intro     本文作者……                 the author bio, printed bold as 机器之心 AIxiv does
  @image     assets/x.png | 图注 | 0.8 | 0.9   path | caption | PDF width | .docx width
                                        as a fraction of the text width
  @paper     论文题目：……                consecutive @paper lines form one bulleted block
  ## 小节标题
  ### 小节内的小标题
  1、**一句加粗的总起** 后面的正文        numbered item, hanging indent
  一个普通段落，可以有 **加粗**，上标 10^{-20}，下标 Z_{q}
  @end                                  closing mark, — 完 —

    build_article.py <source.txt> <out-stem>     writes <out-stem>.docx and .pdf

Both outputs are checked after they are written: every image must be
embedded in the .docx, and the PDF must contain CJK codepoints and report no
missing glyphs, because a PDF built without a CJK font still exits 0 and
still has pages, it just shows tofu.
"""
import os
import pathlib
import re
import shutil
import subprocess
import sys
import tempfile
import urllib.request
import zlib

CJK_FONT = "Noto Sans SC"
FONT_URLS = {
    "NotoSansSC-Regular.otf": "https://github.com/notofonts/noto-cjk/raw/main/Sans/SubsetOTF/SC/NotoSansSC-Regular.otf",
    "NotoSansSC-Bold.otf": "https://github.com/notofonts/noto-cjk/raw/main/Sans/SubsetOTF/SC/NotoSansSC-Bold.otf",
}
ACCENT = (0xEA, 0x58, 0x0C)      # heading accent, the paper's quantum orange
GREY = (0x6B, 0x72, 0x80)
INTRO_BG = "F3F4F6"


# --------------------------------------------------------------------------- font
def ensure_cjk_font():
    """Install Noto Sans SC if fontconfig cannot see it. Cached on the
    persistent mount when DEVTOOLS_HOME is set, because ~/.fonts lives in the
    container filesystem and is gone after a relaunch."""
    have = subprocess.run(["fc-list", ":lang=zh", "family"], capture_output=True, text=True).stdout
    if CJK_FONT in have:
        return
    cache = pathlib.Path(os.environ.get("DEVTOOLS_HOME", str(pathlib.Path.home() / ".local/share"))) / "fonts/noto-cjk"
    cache.mkdir(parents=True, exist_ok=True)
    home_fonts = pathlib.Path.home() / ".fonts"
    home_fonts.mkdir(exist_ok=True)
    for name, url in FONT_URLS.items():
        src = cache / name
        if not src.exists():
            print(f"downloading {name}")
            urllib.request.urlretrieve(url, src)
        shutil.copy2(src, home_fonts / name)
    subprocess.run(["fc-cache", "-f"], check=True, capture_output=True)
    have = subprocess.run(["fc-list", ":lang=zh", "family"], capture_output=True, text=True).stdout
    if CJK_FONT not in have:
        raise SystemExit(f"{CJK_FONT} still not visible to fontconfig after install")


# --------------------------------------------------------------------------- parse
def parse(src):
    base = pathlib.Path(src).resolve().parent
    items, paper = [], []

    def flush_paper():
        if paper:
            items.append(("paper", list(paper)))
            paper.clear()

    for raw in pathlib.Path(src).read_text().split("\n"):
        line = raw.strip()
        if not line or line.startswith("%"):
            continue
        if line.startswith("@paper"):
            paper.append(line[len("@paper"):].strip())
            continue
        flush_paper()
        if line.startswith("@title"):
            items.append(("title", line[6:].strip()))
        elif line.startswith("@subtitle"):
            items.append(("subtitle", line[9:].strip().lstrip("—-－ ")))
        elif line.startswith("@intro"):
            items.append(("intro", line[6:].strip()))
        elif line.startswith("@image"):
            parts = [p.strip() for p in line[6:].split("|")]
            path = (base / parts[0]).resolve()
            if not path.exists():
                raise SystemExit(f"image not found: {parts[0]}")
            cap = parts[1] if len(parts) > 1 else ""
            width = float(parts[2]) if len(parts) > 2 and parts[2] else 1.0
            # optional 4th field: a separate width for the .docx. The PDF is a
            # paged preview and may need an image smaller to fit a page; the
            # .docx goes into the WeChat editor, which has no pages.
            dwidth = float(parts[3]) if len(parts) > 3 and parts[3] else width
            items.append(("image", (str(path), cap, width, dwidth)))
        elif line == "@end":
            items.append(("end", None))
        elif line.startswith("### "):
            items.append(("h3", line[4:].strip()))
        elif line.startswith("## "):
            items.append(("h", line[3:].strip()))
        elif re.match(r"^\d+、", line):
            items.append(("item", line))
        else:
            items.append(("p", line))
    flush_paper()
    kinds = [k for k, _ in items]
    for need in ("title", "subtitle"):
        if kinds.count(need) != 1:
            raise SystemExit(f"source needs exactly one @{need}, found {kinds.count(need)}")
    return items


BOLD = re.compile(r"\*\*(.+?)\*\*")
SCRIPT = re.compile(r"([\^_])\{([^{}]*)\}")


def segments(text):
    """Split text into (content, bold, script) runs. **bold** spans may contain
    ^{superscript} and _{subscript}, which is how 4×10^{-20} and Θ(n^{2}) are
    written in a source: Unicode superscript digits are missing from Noto Sans
    SC, and a bare caret would print literally in both formats."""
    out, pos = [], 0
    parts = []
    for m in BOLD.finditer(text):
        if m.start() > pos:
            parts.append((text[pos:m.start()], False))
        parts.append((m.group(1), True))
        pos = m.end()
    if pos < len(text):
        parts.append((text[pos:], False))
    for chunk, bold in parts:
        p = 0
        for m in SCRIPT.finditer(chunk):
            if m.start() > p:
                out.append((chunk[p:m.start()], bold, None))
            out.append((m.group(2), bold, "sup" if m.group(1) == "^" else "sub"))
            p = m.end()
        if p < len(chunk):
            out.append((chunk[p:], bold, None))
    return out


# --------------------------------------------------------------------------- docx
def build_docx(items, out):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.5)
    text_w = sec.page_width - sec.left_margin - sec.right_margin

    normal = doc.styles["Normal"]
    normal.font.name = CJK_FONT
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)

    def run(p, text, size=11, bold=False, color=None):
        r = p.add_run(text)
        r.font.size, r.font.bold, r.font.name = Pt(size), bold, CJK_FONT
        rf = r._element.rPr.rFonts
        for k in ("w:eastAsia", "w:ascii", "w:hAnsi"):
            rf.set(qn(k), CJK_FONT)
        if color:
            r.font.color.rgb = RGBColor(*color)
        return r

    def rich(p, text, size=11, color=None):
        for content, bold, script in segments(text):
            r = run(p, content, size, bold=bold, color=color)
            if script == "sup":
                r.font.superscript = True
            elif script == "sub":
                r.font.subscript = True

    def para(before=0, after=8, align=None, spacing=1.75):
        p = doc.add_paragraph()
        f = p.paragraph_format
        f.space_before, f.space_after, f.line_spacing = Pt(before), Pt(after), spacing
        if align is not None:
            p.alignment = align
        return p

    def shade(p, fill):
        ppr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
        ppr.append(shd)

    for kind, val in items:
        if kind == "title":
            rich(para(after=4, spacing=1.3), val, size=20)
            for r in doc.paragraphs[-1].runs:
                r.font.bold = True
        elif kind == "subtitle":
            run(para(after=14, spacing=1.3), "——" + val, size=14, color=GREY)
        elif kind == "intro":
            # 机器之心 AIxiv prints the author bio as one bold paragraph
            # (published OSCAR article, 2025-02-03), so it is bold, not boxed.
            p = para(before=2, after=12, spacing=1.6)
            rich(p, "**" + val.replace("**", "") + "**", size=10.5)
        elif kind == "image":
            path, cap, _, dwidth = val
            p = para(before=6, after=2 if cap else 10, align=WD_ALIGN_PARAGRAPH.CENTER, spacing=1.0)
            p.add_run().add_picture(path, width=int(text_w * dwidth))
            if cap:
                run(para(after=12, align=WD_ALIGN_PARAGRAPH.CENTER, spacing=1.3), cap, size=9.5, color=GREY)
        elif kind == "paper":
            for i, line in enumerate(val):
                p = para(after=2 if i < len(val) - 1 else 14, spacing=1.4)
                p.paragraph_format.left_indent = Cm(0.5)
                p.paragraph_format.first_line_indent = Cm(-0.4)
                rich(p, "• " + line, size=10.5)
        elif kind == "h":
            p = para(before=18, after=8, spacing=1.4)
            run(p, val, size=15, bold=True, color=ACCENT)
        elif kind == "h3":
            p = para(before=10, after=6, spacing=1.4)
            run(p, val, size=12.5, bold=True)
        elif kind == "item":
            p = para(after=8)
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.first_line_indent = Cm(-0.6)
            rich(p, val)
        elif kind == "p":
            rich(para(after=10), val)
        elif kind == "end":
            run(para(before=16, after=0, align=WD_ALIGN_PARAGRAPH.CENTER), "— 完 —", size=10, color=GREY)
    doc.save(out)

    # verify: every image went in
    from docx import Document as Reopen
    want = sum(1 for k, _ in items if k == "image")
    got = len(Reopen(out).inline_shapes)
    if got != want:
        raise SystemExit(f"docx has {got} images, source has {want}")
    return want


# --------------------------------------------------------------------------- pdf
TEX_PRE = r"""\documentclass[11pt]{article}
\usepackage[a4paper,margin=2.5cm]{geometry}
\usepackage{xeCJK}
\usepackage{graphicx}
\usepackage[dvipsnames]{xcolor}
\usepackage[most]{tcolorbox}
\usepackage[hidelinks]{hyperref}
\usepackage{needspace}
\setCJKmainfont{Noto Sans SC}[BoldFont={Noto Sans SC Bold}]
\setmainfont{DejaVu Sans}
\urlstyle{same}
\definecolor{accent}{RGB}{234,88,12}
\definecolor{muted}{RGB}{107,114,128}
\definecolor{introbg}{HTML}{F3F4F6}
\linespread{1.6}
\setlength{\parindent}{0pt}
\setlength{\parskip}{8pt}
\pagestyle{plain}
% English words inside a Chinese article must not hyphenate: a paper title
% broken as "mod-els" reads as a typo. Lines go ragged instead.
\hyphenpenalty=10000 \exhyphenpenalty=10000 \sloppy
\begin{document}
"""


def tex_escape(s):
    for a, b in (("\\", r"\textbackslash{}"), ("&", r"\&"), ("%", r"\%"), ("$", r"\$"),
                 ("#", r"\#"), ("_", r"\_"), ("{", r"\{"), ("}", r"\}"),
                 ("~", r"\textasciitilde{}"), ("^", r"\textasciicircum{}")):
        s = s.replace(a, b)
    return s


def tex_rich(s):
    out = []
    for content, bold, script in segments(s):
        t = tex_escape(content)
        t = re.sub(r"(https?://[^\s，。；）】]+)", lambda m: r"\url{" + m.group(1).replace(r"\_", "_") + "}", t)
        if script == "sup":
            t = r"\textsuperscript{" + t + "}"
        elif script == "sub":
            t = r"\textsubscript{" + t + "}"
        if bold:
            t = r"\textbf{" + t + "}"
        out.append(t)
    return "".join(out)


def build_pdf(items, out):
    o = [TEX_PRE]
    for kind, val in items:
        if kind == "title":
            o.append(r"{\linespread{1}\fontsize{20}{28}\selectfont\bfseries " + tex_rich(val) + r"\par}\vspace{4pt}")
        elif kind == "subtitle":
            o.append(r"{\linespread{1}\fontsize{14}{20}\selectfont\color{muted}——" + tex_rich(val) + r"\par}\vspace{8pt}")
        elif kind == "intro":
            o.append(r"{\fontsize{10.5}{17}\selectfont\bfseries " + tex_rich(val.replace("**", "")) + r"\par}\vspace{4pt}")
        elif kind == "image":
            path, cap, width, _ = val
            # image and caption in one minipage so a page break cannot split them
            o.append(r"\begin{center}\begin{minipage}{\linewidth}\centering\includegraphics[width=" + f"{width:.2f}"
                     + r"\linewidth,height=0.34\textheight,keepaspectratio]{" + path + "}")
            if cap:
                o.append(r"\\[4pt]{\linespread{1}\fontsize{9.5}{13}\selectfont\color{muted}" + tex_rich(cap) + "}")
            o.append(r"\end{minipage}\end{center}")
        elif kind == "paper":
            # ragged right: an unbreakable English paper title otherwise forces the justifier
            # to letter-space the Chinese label ("论 文 题 目")
            o.append(r"\begin{itemize}\raggedright\setlength{\itemsep}{0pt}\fontsize{10.5}{16}\selectfont "
                     + " ".join(r"\item " + tex_rich(l) for l in val) + r"\end{itemize}")
        elif kind == "h":
            # a heading must not be stranded at the foot of a page: keep room for
            # it and whatever comes next, which is usually a figure
            o.append(r"\needspace{0.42\textheight}\vspace{8pt}{\linespread{1}\fontsize{15}{22}\selectfont\bfseries\color{accent}" + tex_rich(val) + r"\par}\nopagebreak")
        elif kind == "h3":
            o.append(r"\vspace{4pt}{\linespread{1}\fontsize{12.5}{18}\selectfont\bfseries " + tex_rich(val) + r"\par}")
        elif kind == "item":
            o.append(r"\hangindent=1.6em\hangafter=1 " + tex_rich(val) + r"\par")
        elif kind == "p":
            o.append(tex_rich(val) + r"\par")
        elif kind == "end":
            o.append(r"\vspace{12pt}\begin{center}{\fontsize{10}{14}\selectfont\color{muted}— 完 —}\end{center}")
    o.append(r"\end{document}")

    with tempfile.TemporaryDirectory() as td:
        tex = pathlib.Path(td) / "article.tex"
        tex.write_text("\n".join(o) + "\n")
        for _ in range(2):
            r = subprocess.run(["xelatex", "-interaction=nonstopmode", "-halt-on-error", tex.name],
                               cwd=td, capture_output=True, text=True)
            if r.returncode != 0:
                log = (pathlib.Path(td) / "article.log").read_text(errors="ignore")
                err = re.findall(r"^!.*$", log, re.M)
                raise SystemExit("xelatex failed: " + "; ".join(err[:3]))
        log = (pathlib.Path(td) / "article.log").read_text(errors="ignore")
        missing = log.count("Missing character")
        if missing:
            raise SystemExit(f"PDF has {missing} missing glyphs: the CJK font did not cover the text")
        pages = re.search(r"Output written on .*\((\d+) pages?", log)
        shutil.copy2(pathlib.Path(td) / "article.pdf", out)

    # verify CJK really rendered, from the PDF's own ToUnicode maps
    data = pathlib.Path(out).read_bytes()
    cjk = 0
    for m in re.finditer(rb"stream\r?\n", data):
        end = data.find(b"endstream", m.end())
        try:
            raw = zlib.decompress(data[m.end():end])
        except Exception:
            continue
        if b"beginbf" in raw:
            cjk += sum(1 for t in re.findall(rb"<[0-9A-Fa-f]{4}>\s*<([0-9A-Fa-f]{4})", raw)
                       if 0x4E00 <= int(t, 16) <= 0x9FFF)
    if cjk == 0:
        raise SystemExit("PDF contains no CJK codepoints: Chinese did not render")
    return int(pages.group(1)) if pages else None, cjk


def main(argv):
    if len(argv) != 3:
        raise SystemExit(__doc__)
    src, stem = argv[1], argv[2]
    ensure_cjk_font()
    items = parse(src)
    n_img = build_docx(items, stem + ".docx")
    pages, cjk = build_pdf(items, stem + ".pdf")
    print(f"{stem}.docx  {n_img} images embedded")
    print(f"{stem}.pdf   {pages} pages, {cjk} CJK glyphs mapped, 0 missing")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
